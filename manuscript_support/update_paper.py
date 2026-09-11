from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches


SOURCE = Path(
    "/Users/pangtai/Desktop/paper_latex/pikan/"
    "ver6.9-Tai_2nd_paper_MF-L3CR_no_third_level_headings.docx"
)
OUT_DIR = Path("/Users/pangtai/Desktop/Experiment/plates with holes/test090701")
OUTPUT = OUT_DIR / "ver7.0-Tai_2nd_paper_MF-L3CR_fixed_three_step.docx"
FIGURE = OUT_DIR / "figure7_fixed_three_step_matched_budget.png"
FIGURE_PDF = OUT_DIR / "figure7_fixed_three_step_matched_budget.pdf"


def replace_paragraph(paragraph, text: str) -> None:
    """Replace text while retaining the paragraph style and first-run formatting."""
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = copy.deepcopy(paragraph.runs[0]._r.rPr)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    replace_paragraph(paragraph, text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def format_inline_orders(paragraph) -> None:
    """Use native Word superscript runs for the two inline asymptotic orders."""
    source = paragraph.text
    source = source.replace("O(\u03b5\u207b\u00b3\u1d1f\u00b2)", "[[ORDER]]")
    source = source.replace("10\u207b\u00b9\u2070", "[[TEN]]")
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = copy.deepcopy(paragraph.runs[0]._r.rPr)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    def add(text: str, superscript: bool = False) -> None:
        run = paragraph.add_run(text)
        if run_properties is not None:
            run._r.insert(0, copy.deepcopy(run_properties))
        run.font.superscript = superscript

    for part in re.split(r"(\[\[ORDER\]\]|\[\[TEN\]\])", source):
        if part == "[[ORDER]]":
            add("O(\u03b5")
            add("\u22123/2", superscript=True)
            add(")")
        elif part == "[[TEN]]":
            add("10")
            add("\u221210", superscript=True)
        elif part:
            add(part)


def make_figure() -> None:
    import matplotlib.pyplot as plt
    import numpy as np
    stages = ["P5", "P6"]
    methods = ["Resumed AdamW", "L-BFGS", "MF-L3CR"]
    values = np.array(
        [
            [0.2030097349, 0.2965925281, 0.1562940873],
            [0.2108883497, 0.3354309573, 0.1523654457],
        ]
    )
    colors = ["#2f7fb8", "#e98b2a", "#2c9a43"]
    x = np.arange(len(stages))
    width = 0.22

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.linewidth": 0.8,
            "xtick.major.width": 0.8,
            "ytick.major.width": 0.8,
        }
    )
    fig, ax = plt.subplots(figsize=(7.2, 3.5), constrained_layout=True)
    for j, (method, color) in enumerate(zip(methods, colors)):
        bars = ax.bar(x + (j - 1) * width, values[:, j], width, label=method, color=color)
        for bar, value in zip(bars, values[:, j]):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                value + 0.008,
                f"{value:.3f}",
                ha="center",
                va="bottom",
                fontsize=8,
            )
    ax.set_ylabel(r"Mean terminal gradient norm $\|\nabla F\|_2$")
    ax.set_xlabel("Parameter stage")
    ax.set_xticks(x, stages)
    ax.set_ylim(0, 0.39)
    ax.grid(axis="y", color="#d9d9d9", linewidth=0.6, alpha=0.8)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(ncol=3, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.14))
    fig.savefig(FIGURE, dpi=600, bbox_inches="tight", facecolor="white")
    fig.savefig(FIGURE_PDF, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if not FIGURE.exists():
        make_figure()
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    replacements = {
        7: (
            "BAM-KAN is developed for finite-element-supervised prediction of layer-resolved maximum "
            "principal stress S1 and displacement magnitude U in reinforced concrete slabs with openings. "
            "Eight response-independent physical/geometric channels and three normalized coordinate channels "
            "are embedded separately and fused through fine- and coarse-scale KAN pathways. The database "
            "contains 186 Abaqus-derived cases, with 137 for training, 30 for validation, and 19 for testing. "
            "Across five seeds, BAM-KAN reduces the mean Opening-band error by 14.87% relative to the "
            "parameter-matched CNN and by 14.91% relative to U-Net, and gives the lowest mean Balanced error. "
            "A second-stage matrix-free \u2113\u2083-cubic regularization method (MF-L3CR) is introduced for checkpoint "
            "refinement. Its separable cubic term admits coordinatewise inner updates, while Hessian-vector "
            "products provide curvature information without forming the Hessian. The full-space analysis yields "
            "O(\u03b5\u207b\u00b3\u1d1f\u00b2) first-order evaluation complexity. In a 20-parameter deterministic audit, "
            "MF-L3CR reaches \u2016\u2207F\u2016\u2082 \u2264 10\u207b\u00b9\u2070 for all five checkpoints. Full-data HVP checks remain "
            "accurate up to all 86,653 trainable parameters. All five P6 runs complete three accepted "
            "full-parameter steps. Under L3CR-derived wall-clock budgets, MF-L3CR reduces the mean terminal "
            "training objective by 2.03% at both P5 and P6 relative to resumed AdamW. L-BFGS attains lower "
            "terminal objectives, so the full-parameter evidence is limited to training refinement and "
            "stationarity behavior."
        ),
        227: (
            "Each HVP is evaluated by automatic differentiation, and micro-batch accumulation forms the "
            "full-data operator without assembling H\u2c7c. This mode is used at P5 and P6, with 44,149 and 86,653 "
            "trainable parameters. In the scalability experiment, each inner solve is capped at three HVPs. "
            "A step that passes the outer descent and ratio tests without satisfying the nominal inner residual "
            "tolerance is recorded as an inexact descent step."
        ),
        237: (
            "(iv) Fixed-step and matched-budget comparison. Every P5 and P6 MF-L3CR run is continued until "
            "three accepted full-parameter steps. Resumed AdamW and L-BFGS start from the same checkpoint. "
            "Their per-seed budgets are set by the corresponding MF-L3CR three-step runtime. A method stops "
            "before launching a full-data operation that cannot finish within the remaining budget."
        ),
        239: (
            "For statistical reporting, Global, Opening-band, and Balanced scores follow Section 2.7. Global "
            "pools all 19 held-out cases. Opening-band uses the three eligible opening cases. Reported field "
            "errors are mean \u00b1 standard deviation over five seeds. Matched-budget optimizer differences are "
            "paired by checkpoint and summarized by 95% paired bootstrap intervals."
        ),
        258: (
            "The trainable dimension increases from 44,149 to 86,653, while one full-data HVP increases from "
            "147.5 to 156.7 s (6.2%). Table 6 reports the fixed-three-step, full-parameter P6 correction from the "
            "five matched validation-selected checkpoints."
        ),
        260: (
            "Every P5 and P6 run completes three accepted steps. All 30 accepted steps strictly decrease F, "
            "and the minimum acceptance ratio is \u03c1 = 0.9614. At P6, all five terminal gradient norms are below "
            "their initial values. Seed 11 reduces F from 0.154129 to 0.147307 and \u2016g\u2016\u2082 from 0.542833 to "
            "0.114482. The accepted inner steps are inexact because the three-HVP cap is reached before the "
            "nominal inner residual tolerance."
        ),
        261: (
            "For the matched-budget comparison, resumed AdamW, L-BFGS, and MF-L3CR are initialized from "
            "identical checkpoints. The per-seed budget is the runtime of the corresponding three-step MF-L3CR "
            "run. Table 7 reports five-seed means."
        ),
        263: (
            "Relative to resumed AdamW, MF-L3CR lowers the mean terminal objective by 2.03% at both P5 and P6. "
            "The paired differences MF-L3CR minus resumed AdamW are -0.003120 at P5 (95% CI [-0.003926, "
            "-0.002491]) and -0.003121 at P6 (95% CI [-0.003976, -0.002415]); all five paired differences are "
            "negative at each stage. The mean terminal gradient norm decreases by 23.01% at P5 and 27.75% at "
            "P6, but only four of five checkpoints agree and both 95% intervals cross zero. The trajectory-minimum "
            "gradient norm is lower for all five checkpoints, with paired intervals [-0.124814, -0.027205] at P5 "
            "and [-0.144051, -0.026349] at P6."
        ),
        264: (
            "L-BFGS attains a lower terminal objective than MF-L3CR for every checkpoint at both stages. "
            "Conversely, MF-L3CR attains a lower terminal gradient norm than L-BFGS for every checkpoint. Thus, "
            "the result supports a conditional stationarity advantage under the tested budgets, not objective "
            "dominance over L-BFGS. Because a baseline may stop before one indivisible full-data operation, mean "
            "runtime mismatches remain below 5.0% at P5 and 4.8% at P6."
        ),
        266: (
            "Figure 7. Mean terminal gradient norms at P5 and P6 under matched L3CR-derived wall-clock budgets."
        ),
        277: (
            "BAM-KAN attains its largest accuracy gain in the opening-adjacent region. The advantage persists "
            "for band radii from one to four voxels. Separate physical/geometry and coordinate embeddings retain "
            "case information and spatial location, while the fine and coarse pathways combine local resolution "
            "with slab-scale context. MF-L3CR reaches 10\u207b\u00b9\u2070 stationarity with the lowest backpropagation work "
            "in the 20-parameter deterministic audit. Full-data HVP checks remain accurate at P5 and P6, and all "
            "five P6 checkpoints complete three accepted full-parameter steps. Under L3CR-derived budgets, "
            "terminal objectives are lower than resumed AdamW for all five checkpoints. The terminal-gradient "
            "intervals nevertheless cross zero, and L-BFGS attains lower terminal objectives. The full-parameter "
            "evidence therefore supports conditional training refinement and stationarity behavior, not general "
            "wall-clock or generalization superiority. The active-subspace mode separately lowers the three "
            "held-out error measures by small amounts."
        ),
        280: (
            "MF-L3CR adds matrix-free curvature refinement to validation-selected checkpoints. The full-space "
            "analysis gives O(\u03b5\u207b\u00b3\u1d1f\u00b2) first-order evaluation complexity. Numerically, MF-L3CR reaches "
            "\u2016\u2207F\u2016\u2082 \u2264 10\u207b\u00b9\u2070 in all five 20-parameter audits. At P6, all five validation-selected "
            "checkpoints complete three accepted steps in the full 86,653-parameter space. Under L3CR-derived "
            "wall-clock budgets, MF-L3CR reduces the terminal objective by 2.03% relative to resumed AdamW at "
            "both P5 and P6. Its mean terminal gradient norms are 23.01% and 27.75% lower, respectively, although "
            "the paired intervals cross zero. L-BFGS gives lower terminal objectives. Hence the full-parameter "
            "result establishes executable inexact descent and conditional stationarity improvement, but not a "
            "general optimization or test-set advantage. The separate held-out active-subspace experiment gives "
            "small decreases in Global, Opening-band, and Balanced means."
        ),
    }

    for index, text in replacements.items():
        replace_paragraph(doc.paragraphs[index], text)
    format_inline_orders(doc.paragraphs[7])
    format_inline_orders(doc.paragraphs[280])

    replace_paragraph(
        doc.paragraphs[259],
        "Table 6. Fixed-three-step full-parameter MF-L3CR correction at P6 for the five validation-selected checkpoints.",
    )
    replace_paragraph(
        doc.paragraphs[262],
        "Table 7. Mean terminal values under matched L3CR-derived wall-clock budgets.",
    )

    set_cell_text(
        doc.tables[1].rows[8].cells[1],
        "P5: 44,149 parameters; P6: 86,653 parameters; P\u2c7c = I; three accepted steps per seed",
    )

    p6_rows = [
        ["Seed", "Accepted", "Initial F", "Terminal F", "Initial \u2016g\u2016\u2082", "Terminal \u2016g\u2016\u2082", "Time (s)"],
        ["11", "3", "0.154129", "0.147307", "0.542833", "0.114482", "1694.5"],
        ["23", "3", "0.139322", "0.135980", "0.241511", "0.131400", "1686.3"],
        ["37", "3", "0.200403", "0.181395", "0.730453", "0.164720", "1678.7"],
        ["51", "3", "0.143884", "0.140808", "0.250169", "0.120772", "1702.6"],
        ["73", "3", "0.149621", "0.146487", "0.234538", "0.230454", "1689.8"],
    ]
    for row, values in zip(doc.tables[5].rows, p6_rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)

    matched_rows = [
        ["Stage", "Method", "Terminal F", "Terminal \u2016g\u2016\u2082", "Minimum \u2016g\u2016\u2082", "Time (s)"],
        ["P5", "Resumed AdamW", "0.153479", "0.203010", "0.203010", "1532.4"],
        ["P5", "L-BFGS", "0.129537", "0.296593", "0.172871", "1478.9"],
        ["P5", "MF-L3CR", "0.150359", "0.156294", "0.136602", "1556.6"],
        ["P6", "Resumed AdamW", "0.153516", "0.210888", "0.210888", "1611.6"],
        ["P6", "L-BFGS", "0.129003", "0.335431", "0.176366", "1609.6"],
        ["P6", "MF-L3CR", "0.150395", "0.152365", "0.137641", "1690.4"],
    ]
    for row, values in zip(doc.tables[6].rows, matched_rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)

    figure_paragraph = doc.paragraphs[265]
    for child in list(figure_paragraph._p):
        if child.tag != qn("w:pPr"):
            figure_paragraph._p.remove(child)
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.add_run().add_picture(str(FIGURE), width=Inches(6.10))

    doc.core_properties.title = "BAM-KAN and MF-L3CR with fixed-three-step full-parameter experiments"
    doc.core_properties.subject = "Revised P5/P6 MF-L3CR experiment"
    doc.save(OUTPUT)

    check = Document(OUTPUT)
    corpus = "\n".join(p.text for p in check.paragraphs)
    corpus += "\n" + "\n".join(
        cell.text for table in check.tables for row in table.rows for cell in row.cells
    )
    forbidden = ["47.93%", "2.50%", "0.152344", "822.8", "the remaining four checkpoints each accept one"]
    remaining = [token for token in forbidden if token in corpus]
    if remaining:
        raise RuntimeError(f"Stale experiment values remain: {remaining}")
    if len(check.inline_shapes) != 7:
        raise RuntimeError(f"Unexpected inline-shape count: {len(check.inline_shapes)}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
