#!/usr/bin/env python3
"""Build a source-grounded journal-style Chinese L3CR experiment report."""

from __future__ import annotations

import json
import subprocess
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "L3CR_训练高精度与反向传播等价量效率优势_数值实验报告.docx"
INK = "172B3A"
ACCENT = "155E52"
MUTED = "546573"
GRID = "CBD3D9"
HEADER_FILL = "EEF3F1"
NOTE_FILL = "F1F6F4"
DOCUMENT_FONT = "Arial Unicode MS"
MATH_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/math"
MATH_EXPRESSIONS = {
    1: r"\widehat y_i^{s}=(x_i^{s})^{\mathsf T}\theta^{s},\quad "
       r"\widehat y_i^{u}=(x_i^{u})^{\mathsf T}\theta^{u},\quad "
       r"\theta=((\theta^{s})^{\mathsf T},(\theta^{u})^{\mathsf T})^{\mathsf T}\in\mathbb R^{20}",
    2: r"F(\theta)=\frac{1}{2N}\sum_{i=1}^{N}\left["
       r"\big((x_i^{s})^{\mathsf T}\theta^{s}-y_i^{s}\big)^2+"
       r"\big((x_i^{u})^{\mathsf T}\theta^{u}-y_i^{u}\big)^2\right]",
    3: r"H=\frac1N\operatorname{diag}\!\left((X^{s})^{\mathsf T}X^{s},"
       r"(X^{u})^{\mathsf T}X^{u}\right),\quad "
       r"q=\frac1N\left((X^{s})^{\mathsf T}y^{s},(X^{u})^{\mathsf T}y^{u}\right)^{\mathsf T}",
    4: r"F(\theta)=\frac12\theta^{\mathsf T}H\theta-q^{\mathsf T}\theta+c,"
       r"\qquad \nabla F(\theta)=H\theta-q",
    5: r"\theta^{\star}=H^{-1}q,\qquad "
       r"F(\theta)-F(\theta^{\star})="
       r"\frac12(\theta-\theta^{\star})^{\mathsf T}H(\theta-\theta^{\star})",
    6: r"m_k(s)=g_k^{\mathsf T}s+\frac12s^{\mathsf T}B_ks+"
       r"\frac{\sigma_k}{6}\sum_{j=1}^{20}|s_j|^3",
    7: r"\nabla m_k(s)=g_k+B_ks+\frac{\sigma_k}{2}|s|\odot s",
    8: r"\operatorname{prox}(v_j)="
       r"\frac{2v_j}{1+\sqrt{1+2\alpha\sigma_k|v_j|}}",
    9: r"\rho_k=\frac{F(\theta_k)-F(\theta_k+s_k)}{m_k(0)-m_k(s_k)}",
    10: r"\sigma_{k+1}=\begin{cases}"
        r"\sigma_k/2,&\rho_k\geq0.75,\\ "
        r"\sigma_k,&0.10\leq\rho_k<0.75,\\ "
        r"2\sigma_k,&\rho_k<0.10.\end{cases}",
    11: r"R_{\mathrm{work}}=1-\frac{67}{383}=82.51\%,\qquad "
        r"\frac{C_{\mathrm{AdamW-0}}}{C_{\mathrm{L3CR}}}="
        r"\frac{383}{67}=5.72",
    12: r"\Delta_i=\log_{10}\lVert\nabla F_{\mathrm{L3CR},i}\rVert_2-"
        r"\log_{10}\lVert\nabla F_{\mathrm{baseline},i}\rVert_2",
}
_MATH_ELEMENTS: dict[int, etree._Element] | None = None


def build_math_elements() -> dict[int, etree._Element]:
    global _MATH_ELEMENTS
    if _MATH_ELEMENTS is not None:
        return _MATH_ELEMENTS

    source = "\n\n".join(f"$$\n{expression}\n$$" for expression in MATH_EXPRESSIONS.values())
    process = subprocess.run(
        ["/usr/local/bin/pandoc", "--from=markdown", "--to=docx"],
        input=source.encode("utf-8"),
        capture_output=True,
        check=True,
    )
    with ZipFile(BytesIO(process.stdout)) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    elements = root.xpath(".//m:oMathPara", namespaces={"m": MATH_NAMESPACE})
    if len(elements) != len(MATH_EXPRESSIONS):
        raise RuntimeError(f"Expected {len(MATH_EXPRESSIONS)} Word equations, received {len(elements)}")
    _MATH_ELEMENTS = dict(zip(MATH_EXPRESSIONS, elements, strict=True))
    return _MATH_ELEMENTS


def set_run_font(run, size: float = 10.5, bold: bool = False, color: str = INK, italic: bool = False) -> None:
    run.font.name = DOCUMENT_FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), DOCUMENT_FONT)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = DOCUMENT_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal_fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_fonts.set(qn(f"w:{key}"), DOCUMENT_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, before, after in (
        ("Title", 19, 2, 8),
        ("Heading 1", 14, 15, 7),
        ("Heading 2", 11.8, 10, 5),
        ("Heading 3", 10.8, 8, 4),
    ):
        style = document.styles[style_name]
        style.font.name = DOCUMENT_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK if style_name == "Title" else ACCENT)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), DOCUMENT_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15

    caption = document.styles["Caption"]
    caption.font.name = DOCUMENT_FONT
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption_fonts = caption._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        caption_fonts.set(qn(f"w:{key}"), DOCUMENT_FONT)

    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    border = title_properties.find(qn("w:pBdr"))
    if border is not None:
        title_properties.remove(border)
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(5)
    caption.paragraph_format.line_spacing = 1.18


def configure_page(document: Document) -> int:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.05)
    section.bottom_margin = Cm(2.05)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("BAM-KAN | L3CR HIGH-ACCURACY REFINEMENT")
    set_run_font(run, 8, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("-  "), 8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    set_run_font(footer.add_run("  -"), 8, color=MUTED)
    usable_emu = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
    return int(round(usable_emu / 635.0))


def add_paragraph(document: Document, text: str, *, bold_lead: str | None = None, after: float = 6, keep: bool = False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        set_run_font(paragraph.add_run(bold_lead), bold=True)
        set_run_font(paragraph.add_run(text[len(bold_lead):]))
    else:
        set_run_font(paragraph.add_run(text))
    return paragraph


def add_math(document: Document, expression: str, number: int | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True

    if number is None or number not in MATH_EXPRESSIONS:
        raise ValueError(f"Missing structured Word equation for display {number!r}: {expression}")
    paragraph._p.append(deepcopy(build_math_elements()[number]))
    if number is not None:
        suffix = paragraph.add_run(f"     ({number})")
        set_run_font(suffix, 9, color=MUTED)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_pr = table._tbl.tblPr
    width = table_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        table_pr.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(widths)))

    indent = table_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table_pr.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")

    layout = table_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    margins = table_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        table_pr.append(margins)
    for edge, value in (("top", 90), ("bottom", 90), ("start", 120), ("end", 120)):
        child = margins.find(qn(f"w:{edge}"))
        if child is None:
            child = OxmlElement(f"w:{edge}")
            margins.append(child)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")

    borders = table_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "bottom", "insideH", "insideV", "left", "right"):
        child = borders.find(qn(f"w:{edge}"))
        if child is None:
            child = OxmlElement(f"w:{edge}")
            borders.append(child)
        if edge in ("top", "bottom"):
            child.set(qn("w:val"), "single")
            child.set(qn("w:sz"), "10")
            child.set(qn("w:color"), ACCENT)
        elif edge == "insideH":
            child.set(qn("w:val"), "single")
            child.set(qn("w:sz"), "3")
            child.set(qn("w:color"), GRID)
        else:
            child.set(qn("w:val"), "nil")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for cell_width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(cell_width))
        grid.append(grid_col)

    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, cell_width in zip(row.cells, widths):
            cell.width = Inches(cell_width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_width_xml = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if cell_width_xml is None:
                cell_width_xml = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(cell_width_xml)
            cell_width_xml.set(qn("w:type"), "dxa")
            cell_width_xml.set(qn("w:w"), str(cell_width))


def table_widths(total: int, weights: list[float]) -> list[int]:
    widths = [int(total * weight / sum(weights)) for weight in weights]
    widths[-1] += total - sum(widths)
    return widths


def add_table(document: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[int], note: str | None = None) -> None:
    label = document.add_paragraph(style="Caption")
    label.paragraph_format.keep_with_next = True
    set_run_font(label.add_run(caption), 9, bold=True, color=INK)

    table = document.add_table(rows=1, cols=len(headers))
    header_row = table.rows[0]
    header_row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell, text in zip(header_row.cells, headers):
        set_cell_shading(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        set_run_font(paragraph.add_run(text), 8.3, bold=True)

    for values in rows:
        cells = table.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, values)):
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.14
            set_run_font(paragraph.add_run(text), 8.2)
    set_table_geometry(table, widths)

    if note:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        set_run_font(paragraph.add_run(note), 8.5, color=MUTED)


def add_figure(document: Document, filename: str, caption: str, width: float = 6.25) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(HERE / filename), width=Inches(width))
    note = document.add_paragraph(style="Caption")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.keep_together = True
    set_run_font(note.add_run(caption), 8.7, color=MUTED)


def add_takeaway(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.left_indent = Cm(0.3)
    paragraph.paragraph_format.right_indent = Cm(0.3)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), NOTE_FILL)
    properties.append(shading)
    set_run_font(paragraph.add_run(text), 10.3, bold=True, color=ACCENT)


def fmt(value: float) -> str:
    return f"{value:.3e}"


def main() -> None:
    facts = json.loads((HERE / "verified_advantage_facts.json").read_text(encoding="utf-8"))
    training = pd.read_csv(HERE / "table_training_precision.csv").set_index("method")
    thresholds = pd.read_csv(HERE / "table_threshold_matrix.csv")
    seeds = pd.read_csv(HERE / "table_seed_pairing.csv")
    paired = pd.read_csv(HERE / "paired_gradient_statistics.csv")
    ablation = pd.read_csv(HERE / "table_l3cr_ablation.csv")
    gram = pd.read_csv(HERE.parent / "test082402" / "bam_gram_audit.csv")

    document = Document()
    configure_styles(document)
    content_width = configure_page(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("L3CR在BAM-KAN输出头局部校正中的\n训练高精度与反向传播等价量效率优势"), 18, bold=True, color=INK)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(subtitle.add_run("基于既有空洞板FEM代理模型的补充数值实验"), 11.5, color=MUTED)

    add_heading(document, "摘要", 1)
    add_paragraph(document, (
        "针对空洞板BAM-KAN代理模型，固定已训练主干，仅对20维双输出头进行高精度局部校正。"
        "采用137个训练工况和2,159,924个材料体素构造与完整训练集均方误差严格等价的二次目标。"
        "在5个独立检查点和500个反向传播等价预算下，L3CR均达到‖∇F‖₂≤10⁻¹⁰；"
        "AdamW-0、标准AdamW及L-BFGS的达标比例分别为3/5、0/5和0/5。"
        f"L3CR达到该阈值的中位工作量为67，而AdamW-0完整轨迹的中位工作量为383，"
        f"相对减少{facts['equivalent_work_reduction_percent']:.2f}%。"
        "该优势属于确定性局部校正中的训练高精度与反向传播等价量效率，不表示测试泛化或墙钟时间优势。"
    ))
    add_takeaway(document, "核心结果：5/5达到10⁻¹⁰；67个反向传播等价量；相对AdamW-0减少82.51%。")

    add_heading(document, "1. 实验目的与证据边界", 1)
    add_paragraph(document, (
        "本节检验L3CR能否在相同初始参数和完整训练目标下，更快逼近训练集一阶驻点。"
        "比较对象为标准AdamW、无权重衰减的AdamW-0及L-BFGS。"
        "评价对象不是整个BAM-KAN的86,653维参数空间，而是冻结主干后由两个输出头共同构成的20维完整可训练空间。"
    ))
    add_paragraph(document, (
        "采用原始梯度范数‖∇F(θ)‖₂衡量训练优化精度，采用梯度计算与Hessian–vector product累计次数衡量反向传播等价工作量。"
        "测试集指标仅用于验证高精度训练是否伴随可观察的预测差异，不参与算法选择、停止或参数调整。"
    ))

    document.add_page_break()
    add_heading(document, "2. 冻结输出头与完整训练集二次目标", 1)
    add_heading(document, "2.1 双输出头参数化", 2)
    add_paragraph(document, (
        "设第i个有效材料体素对应的应力头与位移头冻结特征分别为φᵢˢ、φᵢᵘ∈ℝ⁹。"
        "将截距并入增广特征，记xᵢˢ=(φᵢˢ,1)、xᵢᵘ=(φᵢᵘ,1)。"
        "两个输出头的参数向量分别为θˢ、θᵘ∈ℝ¹⁰，总参数向量为θ∈ℝ²⁰。"
    ))
    add_math(document, "ŷᵢˢ = (xᵢˢ)ᵀθˢ,     ŷᵢᵘ = (xᵢᵘ)ᵀθᵘ,     θ = ((θˢ)ᵀ,(θᵘ)ᵀ)ᵀ ∈ ℝ²⁰", 1)
    add_paragraph(document, "令N为全部训练材料体素数。本实验中N=2,159,924。双输出平均平方损失写为")
    add_math(document, "F(θ) = 1/(2N) Σᵢ₌₁ᴺ [( (xᵢˢ)ᵀθˢ − yᵢˢ )² + ( (xᵢᵘ)ᵀθᵘ − yᵢᵘ )²]", 2)

    add_heading(document, "2.2 Gram矩阵构造与精确最优值", 2)
    add_paragraph(document, "将全部训练工况分批累积，定义两个通道的设计矩阵Xˢ、Xᵘ及目标向量yˢ、yᵘ。")
    add_math(document, "H = (1/N) diag((Xˢ)ᵀXˢ,(Xᵘ)ᵀXᵘ),     q = (1/N)((Xˢ)ᵀyˢ,(Xᵘ)ᵀyᵘ)ᵀ", 3)
    add_math(document, "F(θ) = ½ θᵀHθ − qᵀθ + c,     ∇F(θ) = Hθ − q", 4)
    add_math(document, "θ* = H⁻¹q,     F(θ) − F(θ*) = ½(θ−θ*)ᵀH(θ−θ*)", 5)
    add_paragraph(document, (
        f"5个检查点对应的Hessian均为满秩20。正特征值条件数范围为{gram.hessian_condition_positive.min():.2f}"
        f"至{gram.hessian_condition_positive.max():.2f}。Gram二次目标与直接完整训练损失的最大绝对差为"
        f"{gram.gram_direct_abs_error.max():.3e}。因此，本实验检验的是完整训练集精度，而不是小批量闭包的局部拟合。"
    ))

    document.add_page_break()
    add_heading(document, "3. L3CR模型与公平比较协议", 1)
    add_heading(document, "3.1 全空间ℓ₃三次正则子问题", 2)
    add_paragraph(document, "在当前点θₖ，记gₖ=∇F(θₖ)、Bₖ=∇²F(θₖ)。20维全空间L3CR模型定义为")
    add_math(document, "mₖ(s) = gₖᵀs + ½sᵀBₖs + (σₖ/6) Σⱼ₌₁²⁰ |sⱼ|³", 6)
    add_math(document, "∇mₖ(s) = gₖ + Bₖs + (σₖ/2)|s|⊙s", 7)
    add_paragraph(document, "Hessian通过自动微分HVP逐列构造并显式对称化。近端梯度内层的坐标更新满足")
    add_math(document, "prox(vⱼ) = 2vⱼ / (1 + √(1 + 2ασₖ|vⱼ|))", 8)
    add_math(document, "ρₖ = [F(θₖ)−F(θₖ+sₖ)] / [mₖ(0)−mₖ(sₖ)]", 9)
    add_math(document, "σₖ₊₁ = σₖ/2  (ρₖ≥0.75);   σₖ  (0.10≤ρₖ<0.75);   2σₖ  (ρₖ<0.10)", 10)
    add_paragraph(document, "每个接受步均满足训练目标严格下降。原版L3CR使用10⁻¹²内层模型残差阈值。")

    add_heading(document, "3.2 实验协议", 2)
    protocol_rows = [
        ["训练工况 / 验证 / 测试", "137 / 30 / 19"],
        ["有效材料体素", "2,159,924"],
        ["可训练参数", "20；BAM主干冻结"],
        ["随机种子", "11, 23, 37, 51, 73"],
        ["计算设备与数值精度", "CPU；单线程；float64"],
        ["主比较预算", "500个反向传播等价量"],
        ["补充墙钟预算", "2秒；连续轨迹取预设快照"],
        ["高精度阈值", "10⁻⁶、10⁻⁸、10⁻¹⁰"],
    ]
    add_table(document, "表1  完整训练目标与公平比较设置", ["项目", "设定"], protocol_rows, table_widths(content_width, [1.25, 2.7]))

    document.add_page_break()
    add_heading(document, "4. 训练梯度高精度优势", 1)
    add_paragraph(document, (
        "图1给出5个检查点上训练梯度范数关于反向传播等价预算的中位轨迹。"
        "为避免非单调迭代造成终点误判，曲线显示截至当前预算已经达到的最小原始梯度范数。"
        "阴影表示L3CR和AdamW-0的四分位范围。"
    ))
    add_figure(document, "figure_01_training_gradient_convergence.png", "图1  固定20维输出头的训练梯度收敛。虚线表示10⁻¹⁰精度阈值。", 6.15)

    rows = []
    for name in ("L3CR", "AdamW-0", "AdamW", "L-BFGS"):
        row = training.loc[name]
        rows.append([
            name,
            fmt(row.gradient_norm_mean),
            fmt(row.objective_gap_mean),
            f"{int(row.reached_1e10_within_500)}/5",
            "--" if row.median_equivalents_within_500 < 0 else f"{row.median_equivalents_within_500:.0f}",
        ])
    add_table(
        document,
        "表2  500反向传播等价预算下的训练精度",
        ["方法", "平均‖∇F‖₂", "平均F−F*", "达到10⁻¹⁰", "中位工作量"],
        rows,
        table_widths(content_width, [1.0, 1.45, 1.4, 1.05, 1.2]),
        "注：中位工作量仅在500预算内成功达到10⁻¹⁰的seed中计算；未达到记为--。",
    )
    add_paragraph(document, (
        f"在相同500等价预算下，L3CR的平均训练梯度为{training.loc['L3CR','gradient_norm_mean']:.3e}，"
        f"已接近float64算术极限。L-BFGS为{training.loc['L-BFGS','gradient_norm_mean']:.3e}；"
        f"AdamW-0为{training.loc['AdamW-0','gradient_norm_mean']:.3e}。"
        "应强调，AdamW-0的均值受两个未在预算内收敛的检查点影响，因此同时报告中位数、逐seed结果与达标比例。"
    ))

    document.add_page_break()
    add_heading(document, "5. 反向传播等价量效率", 1)
    add_paragraph(document, (
        "达到给定梯度阈值的工作量直接反映优化路径所消耗的梯度与HVP计算。"
        "图2同时给出完整轨迹的中位阈值工作量，以及500预算内的成功次数。"
        "两项指标分别反映计算路径长度和预算约束下的稳定性。"
    ))
    add_figure(document, "figure_02_threshold_efficiency.png", "图2  达到10⁻¹⁰训练梯度阈值的中位工作量与500预算内成功率。", 6.3)

    threshold_rows = []
    for _, row in thresholds.iterrows():
        threshold_rows.append([
            f"{float(row.threshold):.0e}",
            f"{int(row.L3CR_within_500)}/5；{row.L3CR_median_equivalents:.0f}",
            f"{int(row.AdamW_0_within_500)}/5；{row.AdamW_0_median_equivalents:.0f}",
            f"{int(row.AdamW_within_500)}/5；" + ("--" if row.AdamW_median_equivalents < 0 else f"{row.AdamW_median_equivalents:.0f}"),
            f"{int(row.L_BFGS_within_500)}/5；" + ("--" if row.L_BFGS_median_equivalents < 0 else f"{row.L_BFGS_median_equivalents:.0f}"),
        ])
    add_table(
        document,
        "表3  不同训练梯度阈值下的成功次数与中位工作量",
        ["阈值", "L3CR", "AdamW-0", "AdamW", "L-BFGS"],
        threshold_rows,
        table_widths(content_width, [0.85, 1.2, 1.25, 1.2, 1.2]),
        "注：每格为“500预算内成功次数；完整轨迹的中位阈值工作量”。工作量超过500时仍保留，以说明超预算情况。",
    )
    add_math(document, "Rwork = 1 − 67/383 = 82.51%,     CAdamW-0/CL3CR = 383/67 = 5.72", 11)
    add_paragraph(document, (
        "因此，在20维确定性二次校正问题上，L3CR以更少的反向传播等价量稳定达到高精度驻点。"
        "该5.72倍比值描述的是工作量效率，不应解释为实际CPU运行时间加速比。"
    ))

    document.add_page_break()
    add_heading(document, "6. 逐seed稳定性与配对分析", 1)
    add_figure(document, "figure_03_seedwise_precision.png", "图3  五个检查点在500等价预算下的配对训练梯度范数。", 6.1)
    seed_rows = []
    for _, row in seeds.iterrows():
        adam_cost = "--" if row.AdamW_0_equivalents_to_1e10 < 0 else f"{row.AdamW_0_equivalents_to_1e10:.0f}"
        seed_rows.append([
            str(int(row.seed)),
            fmt(row.L3CR_gradient_norm),
            fmt(row.AdamW_0_gradient_norm),
            fmt(row.L_BFGS_gradient_norm),
            f"{row.L3CR_equivalents_to_1e10:.0f}",
            adam_cost,
        ])
    add_table(
        document,
        "表4  五个检查点的逐seed高精度结果",
        ["Seed", "L3CR梯度", "AdamW-0梯度", "L-BFGS梯度", "L3CR成本", "AdamW-0成本"],
        seed_rows,
        table_widths(content_width, [0.6, 1.15, 1.25, 1.2, 0.95, 1.15]),
        "注：梯度为500预算内最后完整状态；成本为完整轨迹中首次达到10⁻¹⁰的反向传播等价量。",
    )
    add_paragraph(document, (
        "L3CR在seed 11、23、37、51和73上均达到10⁻¹⁰。"
        "其中seed 37和73的Hessian条件数分别为7,823.60和2,260.56，"
        "AdamW-0达到同一精度分别需要1,080和9,916个反向传播等价量，均超过主比较预算。"
    ))

    document.add_page_break()
    add_heading(document, "7. 配对统计与实现一致性", 1)
    add_paragraph(document, "对于每个seed，定义训练梯度的对数尺度配对差：")
    add_math(document, "Δᵢ = log₁₀‖∇FL3CR,i‖₂ − log₁₀‖∇Fbaseline,i‖₂", 12)
    add_paragraph(document, "采用seed级配对bootstrap计算95%区间。负值表示L3CR在相同工作量下达到更低训练梯度。")
    paired_rows = []
    for _, row in paired.iterrows():
        paired_rows.append([
            row.comparison.replace("L3CR versus ", "L3CR vs "),
            f"{row.mean_difference:.2f}",
            f"[{row.ci95_low:.2f}, {row.ci95_high:.2f}]",
            f"{int(row.l3cr_lower_count)}/{int(row.seed_count)}",
        ])
    add_table(
        document,
        "表5  训练梯度对数尺度的seed级配对差",
        ["比较", "均值差", "95% bootstrap区间", "L3CR更低"],
        paired_rows,
        table_widths(content_width, [1.5, 0.9, 1.6, 0.9]),
        "注：仅有5个seed，区间用于描述条件性方向稳定性，不用于宣称普适统计显著性。",
    )

    add_heading(document, "7.1 L3CR实现版本的消融", 2)
    ablation_rows = []
    for _, row in ablation.iterrows():
        ablation_rows.append([
            row.method,
            fmt(row.gradient_norm_mean),
            f"{int(row.reached_1e10_within_500)}/5",
            f"{row.median_equivalents:.0f}",
            f"{row.median_time_seconds:.3f}",
        ])
    add_table(
        document,
        "表6  三种L3CR实现的20维输出头消融",
        ["实现", "平均‖∇F‖₂", "预算内达标", "中位工作量", "中位时间/s"],
        ablation_rows,
        table_widths(content_width, [1.3, 1.25, 1.0, 1.0, 1.05]),
    )
    add_paragraph(document, (
        "三种实现均以67个中位反向传播等价量达到10⁻¹⁰。"
        "因此，核心优势来自20维全空间二阶校正机制，而不是自适应内层容差本身。"
        "在单独的17参数非线性实验中，用户自适应版本出现112个未满足内层forcing条件的接受步；"
        "该问题未发生在本节的二次输出头实验中。"
    ))

    document.add_page_break()
    add_heading(document, "8. 预测结果与结论边界", 1)
    boundary_rows = []
    for name in ("L3CR", "AdamW-0", "AdamW", "L-BFGS"):
        row = training.loc[name]
        boundary_rows.append([
            name,
            f"{row.global_mean:.6f}",
            f"{row.hole_mean:.6f}",
            f"{row.balanced_mean:.6f}",
        ])
    add_table(
        document,
        "表7  500工作量预算下的测试集指标",
        ["方法", "Global", "Opening", "Balanced"],
        boundary_rows,
        table_widths(content_width, [1.1, 1.2, 1.2, 1.2]),
    )
    add_paragraph(document, (
        "各收敛路线的Global、Opening和Balanced在显示精度内基本一致。"
        "因此，本文识别的优势是训练目标的高精度局部求解，不能表述为测试集泛化提升。"
    ))
    add_paragraph(document, (
        f"此外，L3CR达到10⁻¹⁰的中位墙钟时间为{facts['l3cr_median_wall_time']:.3f} s，"
        f"AdamW-0为{facts['adamw0_median_wall_time']:.3f} s。"
        "L3CR具有反向传播等价量优势，但在当前CPU小规模实现中不具有墙钟时间优势。"
        "该结论不能推广至完整86,653参数网络或一般非线性神经网络训练。"
    ))
    add_heading(document, "9. 可直接用于论文的实验结论", 1)
    add_takeaway(document, (
        "在冻结BAM-KAN主干后的20维满秩输出头校正问题中，L3CR在全部5个检查点上"
        "以67个中位反向传播等价量达到10⁻¹⁰训练梯度阈值。相较于AdamW-0的383个中位等价量，"
        "工作量减少82.51%。在500等价预算下，L3CR与AdamW-0的达标率分别为5/5和3/5。"
        "结果支持L3CR在小型确定性二次局部校正中的训练高精度与工作量效率优势。"
    ))
    add_paragraph(document, (
        "数据来源：test082402/bam_snapshot_metrics.csv、bam_diagnostics.csv、"
        "threshold_reach_summary.csv及bam_gram_audit.csv。所有表格和图像均由冻结原始CSV重新计算。"
    ), after=0)

    document.core_properties.title = "L3CR在BAM-KAN局部校正中的训练高精度与反向传播等价量效率"
    document.core_properties.subject = "Source-grounded numerical evidence for full-space L3CR refinement"
    document.core_properties.author = "Numerical Optimization Research"
    document.core_properties.keywords = "L3CR; BAM-KAN; high-accuracy refinement; Hessian-vector products"
    document.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
